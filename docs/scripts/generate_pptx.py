#!/usr/bin/env python3
"""
MaktabaBora PowerPoint Presentation Generator
Creates a professional, high-grade 16:9 widescreen PowerPoint deck (.pptx)
from MaktabaBora_System_Documentation.md with all text, tables, and 11 Horizontal A4 diagrams.
"""

import os
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Paths
REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DOCS_DIR = os.path.join(REPO_ROOT, "docs")
DIAGRAMS_DIR = os.path.join(DOCS_DIR, "diagrams")
OUTPUT_PPTX = os.path.join(DOCS_DIR, "MaktabaBora_System_Presentation.pptx")

# Color Palette (MaktabaBora Professional System Theme)
COLOR_BG_DARK = RGBColor(15, 23, 42)      # Slate 900 (Deep Navy)
COLOR_BG_LIGHT = RGBColor(248, 250, 252)  # Slate 50 (Off-white canvas)
COLOR_CARD_BG = RGBColor(255, 255, 255)   # Pure White
COLOR_CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
COLOR_TEXT_PRIMARY = RGBColor(15, 23, 42) # Slate 900
COLOR_TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
COLOR_TEXT_LIGHT = RGBColor(241, 245, 249) # Slate 100
COLOR_ACCENT_BLUE = RGBColor(2, 132, 199)  # Sky 600
COLOR_ACCENT_CYAN = RGBColor(14, 165, 233) # Sky 500
COLOR_ACCENT_EMERALD = RGBColor(16, 185, 129) # Emerald 500
COLOR_ACCENT_AMBER = RGBColor(245, 158, 11)  # Amber 500
COLOR_CODE_BG = RGBColor(30, 41, 59)      # Slate 800
COLOR_CODE_TEXT = RGBColor(56, 189, 248)   # Sky 400


class MaktabaBoraDeckBuilder:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank_layout = self.prs.slide_layouts[6]
        self.slide_count = 0

    def add_blank_slide(self, is_dark=False):
        self.slide_count += 1
        slide = self.prs.slides.add_slide(self.blank_layout)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG_DARK if is_dark else COLOR_BG_LIGHT
        bg.line.fill.background()
        return slide

    def add_header(self, slide, title, category="MAKTABABORA ARCHITECTURE", is_dark=False):
        # Category Badge
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Segoe UI"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT_CYAN if is_dark else COLOR_ACCENT_BLUE

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.6))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_LIGHT if is_dark else COLOR_TEXT_PRIMARY

        # Subtle Accent Rule
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(2.0), Inches(0.04))
        rule.fill.solid()
        rule.fill.fore_color.rgb = COLOR_ACCENT_CYAN if is_dark else COLOR_ACCENT_BLUE
        rule.line.fill.background()

    def add_card(self, slide, left, top, width, height, title=None, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()

        if title:
            tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = "Segoe UI"
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEXT_PRIMARY

        return card

    def add_diagram_slide(self, title, image_filename, category, takeaways):
        """Creates a high-impact slide featuring a Horizontal A4 diagram on left and executive notes on right."""
        slide = self.add_blank_slide(is_dark=False)
        self.add_header(slide, title, category=category, is_dark=False)

        img_path = os.path.join(DIAGRAMS_DIR, image_filename)
        # Diagram placement: Left = 0.8 in, Top = 1.5 in, Width = 7.6 in, Height = 5.37 in (1.414 ratio)
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, Inches(0.8), Inches(1.5), width=Inches(7.6), height=Inches(5.375))
            # Subtle frame border around picture
            frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(7.6), Inches(5.375))
            frame.fill.background()
            frame.line.color.rgb = COLOR_CARD_BORDER
            frame.line.width = Pt(1)
        else:
            ph = self.add_card(slide, 0.8, 1.5, 7.6, 5.375, title=f"Diagram: {image_filename}")

        # Right Side: Executive Technical Takeaways Card
        right_card = self.add_card(slide, 8.65, 1.5, 3.88, 5.375, title="Technical Specifications")
        tb = slide.shapes.add_textbox(Inches(8.85), Inches(2.1), Inches(3.48), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        for i, item in enumerate(takeaways):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.space_after = Pt(10)
            if isinstance(item, tuple):
                h, b = item
                run1 = p.add_run()
                run1.text = f"{h}: "
                run1.font.name = "Segoe UI"
                run1.font.bold = True
                run1.font.size = Pt(11)
                run1.font.color.rgb = COLOR_ACCENT_BLUE

                run2 = p.add_run()
                run2.text = b
                run2.font.name = "Segoe UI"
                run2.font.size = Pt(11)
                run2.font.color.rgb = COLOR_TEXT_PRIMARY
            else:
                run = p.add_run()
                run.text = f"• {item}"
                run.font.name = "Segoe UI"
                run.font.size = Pt(11)
                run.font.color.rgb = COLOR_TEXT_PRIMARY

    def add_table_slide(self, title, category, headers, rows, col_widths=None):
        slide = self.add_blank_slide(is_dark=False)
        self.add_header(slide, title, category=category, is_dark=False)

        num_rows = len(rows) + 1
        num_cols = len(headers)
        table_width = Inches(11.733)
        table_height = Inches(min(5.2, 0.45 * num_rows))

        table_shape = slide.shapes.add_table(num_rows, num_cols, Inches(0.8), Inches(1.6), table_width, table_height)
        table = table_shape.table

        # Column widths
        if col_widths:
            for c_idx, w in enumerate(col_widths):
                table.columns[c_idx].width = Inches(w)

        # Header Row
        for c_idx, h in enumerate(headers):
            cell = table.cell(0, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_BG_DARK
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.12)
            tf.margin_top = tf.margin_bottom = Inches(0.08)
            p = tf.paragraphs[0]
            p.text = h
            p.font.name = "Segoe UI"
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEXT_LIGHT

        # Data Rows
        for r_idx, row in enumerate(rows):
            is_even = (r_idx % 2 == 0)
            for c_idx, val in enumerate(row):
                cell = table.cell(r_idx + 1, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_CARD_BG if is_even else RGBColor(241, 245, 249)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = Inches(0.12)
                tf.margin_top = tf.margin_bottom = Inches(0.06)
                p = tf.paragraphs[0]
                p.text = str(val)
                p.font.name = "Segoe UI"
                p.font.size = Pt(10)
                p.font.color.rgb = COLOR_TEXT_PRIMARY

    def add_code_slide(self, title, category, code_text, explanation_points, language="PHP"):
        slide = self.add_blank_slide(is_dark=False)
        self.add_header(slide, title, category=category, is_dark=False)

        # Left: Code Container
        code_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(7.2), Inches(5.375))
        code_card.fill.solid()
        code_card.fill.fore_color.rgb = COLOR_CODE_BG
        code_card.line.fill.background()

        # Language pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.65), Inches(1.0), Inches(0.28))
        pill.fill.solid()
        pill.fill.fore_color.rgb = COLOR_ACCENT_BLUE
        pill.line.fill.background()
        p_pill = pill.text_frame.paragraphs[0]
        p_pill.text = language
        p_pill.font.name = "Consolas"
        p_pill.font.size = Pt(9)
        p_pill.font.bold = True
        p_pill.font.color.rgb = COLOR_TEXT_LIGHT

        tb_code = slide.shapes.add_textbox(Inches(1.0), Inches(2.05), Inches(6.8), Inches(4.6))
        tf_code = tb_code.text_frame
        tf_code.word_wrap = True
        tf_code.margin_left = tf_code.margin_top = tf_code.margin_right = tf_code.margin_bottom = 0
        p_c = tf_code.paragraphs[0]
        p_c.text = code_text
        p_c.font.name = "Consolas"
        p_c.font.size = Pt(9.5)
        p_c.font.color.rgb = COLOR_CODE_TEXT

        # Right: Explanation Card
        self.add_card(slide, 8.25, 1.5, 4.28, 5.375, title="Implementation Logic")
        tb_exp = slide.shapes.add_textbox(Inches(8.45), Inches(2.1), Inches(3.88), Inches(4.5))
        tf_exp = tb_exp.text_frame
        tf_exp.word_wrap = True
        tf_exp.margin_left = tf_exp.margin_top = tf_exp.margin_right = tf_exp.margin_bottom = 0

        for i, pt in enumerate(explanation_points):
            p = tf_exp.add_paragraph() if i > 0 else tf_exp.paragraphs[0]
            p.space_after = Pt(12)
            if isinstance(pt, tuple):
                h, b = pt
                run1 = p.add_run()
                run1.text = f"{h}\n"
                run1.font.name = "Segoe UI"
                run1.font.bold = True
                run1.font.size = Pt(11)
                run1.font.color.rgb = COLOR_ACCENT_BLUE

                run2 = p.add_run()
                run2.text = b
                run2.font.name = "Segoe UI"
                run2.font.size = Pt(10.5)
                run2.font.color.rgb = COLOR_TEXT_PRIMARY
            else:
                run = p.add_run()
                run.text = f"• {pt}"
                run.font.name = "Segoe UI"
                run.font.size = Pt(10.5)
                run.font.color.rgb = COLOR_TEXT_PRIMARY

    def build_deck(self):
        print("Starting MaktabaBora PowerPoint generation...")

        # -------------------------------------------------------------
        # Slide 1: Title Slide (Dark Theme)
        # -------------------------------------------------------------
        s1 = self.add_blank_slide(is_dark=True)
        # Accent Glow / Decorative Card
        dec = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.15), Inches(3.8))
        dec.fill.solid()
        dec.fill.fore_color.rgb = COLOR_ACCENT_CYAN
        dec.line.fill.background()

        tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = "ENTERPRISE TECHNICAL ARCHITECTURE & SYSTEMS SPECIFICATION"
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_ACCENT_CYAN
        p0.space_after = Pt(12)

        p1 = tf.add_paragraph()
        p1.text = "MaktabaBora Smart Library Management System"
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(36)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_LIGHT
        p1.space_after = Pt(14)

        p2 = tf.add_paragraph()
        p2.text = "A modern hybrid library ecosystem uniting automated physical circulation, in-browser digital eBook access, cashless Safaricom M-Pesa payments, and grounded AI retrieval (RAG)."
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(15)
        p2.font.color.rgb = RGBColor(148, 163, 184) # Slate 400
        p2.space_after = Pt(28)

        p3 = tf.add_paragraph()
        p3.text = "Core Stack: PHP 8.2+ · Laravel 11.x · React 18 · Tailwind CSS v4 · PostgreSQL 16 · Neon Cloud · OpenAI / Gemini RAG"
        p3.font.name = "Consolas"
        p3.font.size = Pt(11)
        p3.font.color.rgb = COLOR_ACCENT_CYAN

        # -------------------------------------------------------------
        # Slide 2: Executive Summary (Dark Theme)
        # -------------------------------------------------------------
        s2 = self.add_blank_slide(is_dark=True)
        self.add_header(s2, "Executive Summary & Core Mission", category="STRATEGIC OVERVIEW", is_dark=True)

        cards = [
            ("Patron & Digital Services", "Bridges physical book circulation with lifetime digital reading. Features an interactive in-browser EPUB/PDF reader, dynamic multi-tier perk subscriptions, and automated hold waitlists.", COLOR_ACCENT_CYAN),
            ("Staff Circulation Desk", "Eliminates manual paper ledgers with automated barcode loans, dynamic overdue fine accrual (KES 10/day), 1-click Open Library ISBN catalog imports, and fine waiver auditing.", COLOR_ACCENT_EMERALD),
            ("Enterprise AI & Payments", "Features a conversational RAG AI Librarian grounded strictly in catalog records with a 20k daily token limiter, alongside simulated Safaricom Daraja M-Pesa STK push checkout.", COLOR_ACCENT_AMBER)
        ]

        for i, (head, desc, acc) in enumerate(cards):
            c_left = 0.8 + i * 4.0
            card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(1.8), Inches(3.7), Inches(5.0))
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(30, 41, 59)
            card.line.color.rgb = RGBColor(51, 65, 85)

            # Header pill
            hpill = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(c_left + 0.3), Inches(2.2), Inches(0.12), Inches(0.35))
            hpill.fill.solid()
            hpill.fill.fore_color.rgb = acc
            hpill.line.fill.background()

            tb_c = s2.shapes.add_textbox(Inches(c_left + 0.55), Inches(2.15), Inches(2.85), Inches(4.3))
            tf_c = tb_c.text_frame
            tf_c.word_wrap = True
            p_h = tf_c.paragraphs[0]
            p_h.text = head
            p_h.font.name = "Segoe UI"
            p_h.font.size = Pt(16)
            p_h.font.bold = True
            p_h.font.color.rgb = COLOR_TEXT_LIGHT
            p_h.space_after = Pt(14)

            p_d = tf_c.add_paragraph()
            p_d.text = desc
            p_d.font.name = "Segoe UI"
            p_d.font.size = Pt(12)
            p_d.font.color.rgb = RGBColor(203, 213, 225)

        # -------------------------------------------------------------
        # Slide 3: Architecture Agenda (Light Theme)
        # -------------------------------------------------------------
        s3 = self.add_blank_slide(is_dark=False)
        self.add_header(s3, "Architecture & Presentation Agenda", category="TABLE OF CONTENTS", is_dark=False)

        agenda_groups = [
            ("01. FOUNDATIONS", ["Institutional Objectives & Features", "High-Level System Architecture", "Core Technology Matrix", "Multi-Tier Separation of Concerns"]),
            ("02. DATA & STRUCTURE", ["Relational 3NF Database & ER Model", "Comprehensive Schema Dictionary", "Directory Layout & Model Classes", "REST Controllers & Domain Services"]),
            ("03. WORKFLOWS & CODE", ["Auth & JWT Token Lifecycle", "Circulation, Returns & Fines", "Hold Queue & Digital Purchase", "AI RAG Assistant & Code Highlights"]),
            ("04. OPERATIONS & QA", ["API Gateway & Integrations", "Multi-Layer Security Defenses", "Testing Suite (127 Passing Tests)", "Deployment, Roadmap & Conclusion"])
        ]

        for i, (grp_title, items) in enumerate(agenda_groups):
            left = 0.8 + (i % 2) * 5.95
            top = 1.6 + (i // 2) * 2.7
            self.add_card(s3, left, top, 5.75, 2.45, title=grp_title)

            tb_a = s3.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.65), Inches(5.15), Inches(1.6))
            tf_a = tb_a.text_frame
            tf_a.word_wrap = True
            tf_a.margin_left = tf_a.margin_top = tf_a.margin_right = tf_a.margin_bottom = 0

            for j, item in enumerate(items):
                p = tf_a.add_paragraph() if j > 0 else tf_a.paragraphs[0]
                p.text = f"•  {item}"
                p.font.name = "Segoe UI"
                p.font.size = Pt(11)
                p.font.color.rgb = COLOR_TEXT_PRIMARY
                p.space_after = Pt(4)

        # -------------------------------------------------------------
        # Slide 4: Institutional Objectives (Light Theme)
        # -------------------------------------------------------------
        s4 = self.add_blank_slide(is_dark=False)
        self.add_header(s4, "Institutional Objectives & Strategic Vision", category="SECTION 1: OBJECTIVES", is_dark=False)

        objs = [
            ("1. Automated Physical Circulation", "Eliminate paper ledgers via real-time barcode tracking, auto-computed return dates, and dynamic fine accrual."),
            ("2. Hybrid Physical + Digital Lending", "Empower remote patrons with in-browser digital reading (EPUB/PDF), chapter navigation, and study notes."),
            ("3. Cashless Financial Operations", "Direct integration with Safaricom Daraja M-Pesa for instant settlement of passes, digital rentals, and fines (simulated in MVP)."),
            ("4. Grounded AI Catalog Discovery", "Equip patrons with a conversational AI Assistant strictly grounded in real catalog holdings with a 20k token daily cost limiter."),
            ("5. Global Catalog Interoperability", "Rapidly populate collections with 1-click ISBN metadata synchronization from Open Library REST APIs."),
            ("6. Defense-in-Depth & Privacy", "Enforce stateless HMAC-SHA256 JWT tokens, granular RBAC, and absolute patron borrowing privacy.")
        ]

        for i, (title_o, desc_o) in enumerate(objs):
            c_left = 0.8 + (i % 3) * 3.95
            c_top = 1.6 + (i // 3) * 2.75
            card = self.add_card(s4, c_left, c_top, 3.8, 2.5)

            tb_o = s4.shapes.add_textbox(Inches(c_left + 0.25), Inches(c_top + 0.25), Inches(3.3), Inches(2.0))
            tf_o = tb_o.text_frame
            tf_o.word_wrap = True
            tf_o.margin_left = tf_o.margin_top = tf_o.margin_right = tf_o.margin_bottom = 0

            p1 = tf_o.paragraphs[0]
            p1.text = title_o
            p1.font.name = "Segoe UI"
            p1.font.size = Pt(13)
            p1.font.bold = True
            p1.font.color.rgb = COLOR_ACCENT_BLUE
            p1.space_after = Pt(8)

            p2 = tf_o.add_paragraph()
            p2.text = desc_o
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(11)
            p2.font.color.rgb = COLOR_TEXT_PRIMARY

        # -------------------------------------------------------------
        # Slide 5: Key Feature Modules - Patron Services (Light Theme)
        # -------------------------------------------------------------
        s5 = self.add_blank_slide(is_dark=False)
        self.add_header(s5, "Key Features: Patron Services & Hybrid Lending", category="SECTION 1: KEY FEATURES", is_dark=False)

        features_patron = [
            ("Public Catalog & Discovery", [
                "Full-text search by title, author, genre, or ISBN",
                "Real-time physical copy availability indicators",
                "Reference-only restriction warnings for special collections",
                "Foreign currency book pricing with live KES estimates"
            ]),
            ("FIFO Hold Reservation Queue", [
                "Place holds on checked-out titles with dynamic queue rank",
                "Automated allocation when physical copies are returned",
                "Printable and downloadable digital hold receipts",
                "1-click patron hold cancellation"
            ]),
            ("Digital Library & e-Reader", [
                "One-time purchase granting lifetime in-app access",
                "In-browser EPUB & PDF reading without browser plugins",
                "Custom themes (Light, Dark, Sepia) and font scaling",
                "Chapter navigation and local margin notes notepad"
            ]),
            ("Membership Perk Passes", [
                "Configurable tiers: General, Student, Scholar, VIP",
                "Dynamic digital book discounts applied in cart (10%-25%)",
                "Higher physical borrow limits (up to 10 books)",
                "Transparent refund claim filing with staff review"
            ])
        ]

        for i, (title_f, bullets) in enumerate(features_patron):
            c_left = 0.8 + (i % 2) * 5.95
            c_top = 1.6 + (i // 2) * 2.7
            self.add_card(s5, c_left, c_top, 5.75, 2.5, title=title_f)

            tb_f = s5.shapes.add_textbox(Inches(c_left + 0.3), Inches(c_top + 0.65), Inches(5.15), Inches(1.7))
            tf_f = tb_f.text_frame
            tf_f.word_wrap = True
            tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0

            for j, b in enumerate(bullets):
                p = tf_f.add_paragraph() if j > 0 else tf_f.paragraphs[0]
                p.text = f"•  {b}"
                p.font.name = "Segoe UI"
                p.font.size = Pt(11)
                p.font.color.rgb = COLOR_TEXT_PRIMARY
                p.space_after = Pt(3)

        # -------------------------------------------------------------
        # Slide 6: Key Feature Modules - Staff & Admin Core (Light Theme)
        # -------------------------------------------------------------
        s6 = self.add_blank_slide(is_dark=False)
        self.add_header(s6, "Key Features: Circulation Desk & Administration", category="SECTION 1: KEY FEATURES", is_dark=False)

        features_admin = [
            ("Librarian Circulation Desk", [
                "High-speed barcode scanner checkout and return processing",
                "Automated overdue fine accrual at KES 10.00 / day past due",
                "Authority to waive patron fines with mandatory audit logging",
                "Custom physical borrow limit overrides per patron"
            ]),
            ("Open Library ISBN Integration", [
                "1-click catalog import from Open Library REST API",
                "Auto-populates title, author, synopsis, and cover art",
                "Batch copy creation with unique shelf barcode generation",
                "Graceful degradation when external services are unreachable"
            ]),
            ("AI Assistant & Quota Limiter", [
                "Conversational RAG assistant grounded strictly in library stock",
                "Anti-hallucination guardrails and PII protection filters",
                "Daily 20,000 token consumption cost limiter per member",
                "Dynamic OpenAI / Gemini API key rotation and testing"
            ]),
            ("System Admin & Dynamic CRUD", [
                "Global circulation analytics: active loans, overdue, revenue",
                "User account governance with ban toggles and reason auditing",
                "First-time librarian password rotation enforcement",
                "Direct administrative table data management console"
            ])
        ]

        for i, (title_f, bullets) in enumerate(features_admin):
            c_left = 0.8 + (i % 2) * 5.95
            c_top = 1.6 + (i // 2) * 2.7
            self.add_card(s6, c_left, c_top, 5.75, 2.5, title=title_f)

            tb_f = s6.shapes.add_textbox(Inches(c_left + 0.3), Inches(c_top + 0.65), Inches(5.15), Inches(1.7))
            tf_f = tb_f.text_frame
            tf_f.word_wrap = True
            tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0

            for j, b in enumerate(bullets):
                p = tf_f.add_paragraph() if j > 0 else tf_f.paragraphs[0]
                p.text = f"•  {b}"
                p.font.name = "Segoe UI"
                p.font.size = Pt(11)
                p.font.color.rgb = COLOR_TEXT_PRIMARY
                p.space_after = Pt(3)

        # -------------------------------------------------------------
        # Slide 7: High-Level System Architecture (Light Theme)
        # -------------------------------------------------------------
        s7 = self.add_blank_slide(is_dark=False)
        self.add_header(s7, "Multi-Tier Client-Server Architecture", category="SECTION 2: ARCHITECTURE", is_dark=False)

        tiers = [
            ("Client Presentation Layer", "React 18 Single Page Application built with Vite 7. Manages client state via AuthContext and CartContext, with responsive Tailwind CSS v4 styling and Framer Motion micro-animations."),
            ("API Gateway & Security", "Centralized /api/v1 gateway enforcing CORS headers, IP rate limiting, HMAC-SHA256 JWT validation, and RBAC privilege gates before requests reach controllers."),
            ("Domain Services Layer", "Decoupled PHP services bound to strict interface contracts. Handles JWT signing, Daraja M-Pesa STK push, vector embeddings, TF-IDF recommendations, and loan quotas."),
            ("Data Persistence Layer", "PostgreSQL 16 relational database utilizing 3NF normalization. Automatically toggles between Neon Cloud connection pooling and direct migration endpoints.")
        ]

        for i, (t_name, t_desc) in enumerate(tiers):
            t_top = 1.6 + i * 1.35
            self.add_card(s7, 0.8, t_top, 11.733, 1.2)

            # Left number badge
            nbadge = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(t_top + 0.25), Inches(0.4), Inches(0.7))
            nbadge.fill.solid()
            nbadge.fill.fore_color.rgb = COLOR_ACCENT_BLUE
            nbadge.line.fill.background()
            p_nb = nbadge.text_frame.paragraphs[0]
            p_nb.text = f"0{i+1}"
            p_nb.font.name = "Segoe UI"
            p_nb.font.size = Pt(14)
            p_nb.font.bold = True
            p_nb.font.color.rgb = COLOR_TEXT_LIGHT

            tb_t = s7.shapes.add_textbox(Inches(1.6), Inches(t_top + 0.2), Inches(10.7), Inches(0.8))
            tf_t = tb_t.text_frame
            tf_t.word_wrap = True
            tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0

            p_th = tf_t.paragraphs[0]
            p_th.text = t_name
            p_th.font.name = "Segoe UI"
            p_th.font.size = Pt(13)
            p_th.font.bold = True
            p_th.font.color.rgb = COLOR_TEXT_PRIMARY

            p_td = tf_t.add_paragraph()
            p_td.text = t_desc
            p_td.font.name = "Segoe UI"
            p_td.font.size = Pt(10.5)
            p_td.font.color.rgb = COLOR_TEXT_MUTED

        # -------------------------------------------------------------
        # Slide 8: Architectural Layers & Separation of Concerns (Light Theme)
        # -------------------------------------------------------------
        s8 = self.add_blank_slide(is_dark=False)
        self.add_header(s8, "Architectural Separation of Concerns", category="SECTION 2: ARCHITECTURE", is_dark=False)

        layers = [
            ("Providers (`app/Providers`)", "The Setup Crew", "Runs before any request is handled. Binds service interfaces to concrete singletons in the IoC container and defines authorization Gates ('access-admin', 'access-librarian')."),
            ("Middleware (`app/Http/Middleware`)", "The Security Gatekeepers", "Inspects requests at the door. Validates JWT signatures, checks banned status, evaluates borrow quotas, and enforces daily 20k AI token limits before controllers execute."),
            ("Controllers (`app/Http/Controllers`)", "The Traffic Directors", "Thin HTTP adapters. Validates request parameters, calls relevant services, and returns formatted JSON responses (200, 201, 401, 403, 422, 500). Contains zero business logic."),
            ("Services (`app/Services`)", "The Domain Brains", "Pure business logic. Performs cryptographic token hashing, Daraja M-Pesa STK pushes, vector similarity calculations, and Open Library catalog imports independent of HTTP.")
        ]

        for i, (l_name, l_sub, l_body) in enumerate(layers):
            c_left = 0.8 + (i % 2) * 5.95
            c_top = 1.6 + (i // 2) * 2.7
            self.add_card(s8, c_left, c_top, 5.75, 2.5)

            tb_l = s8.shapes.add_textbox(Inches(c_left + 0.3), Inches(c_top + 0.25), Inches(5.15), Inches(2.0))
            tf_l = tb_l.text_frame
            tf_l.word_wrap = True
            tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0

            p1 = tf_l.paragraphs[0]
            p1.text = l_name
            p1.font.name = "Segoe UI"
            p1.font.size = Pt(13)
            p1.font.bold = True
            p1.font.color.rgb = COLOR_TEXT_PRIMARY

            p2 = tf_l.add_paragraph()
            p2.text = f"Role: {l_sub}"
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(10.5)
            p2.font.bold = True
            p2.font.color.rgb = COLOR_ACCENT_BLUE
            p2.space_after = Pt(6)

            p3 = tf_l.add_paragraph()
            p3.text = l_body
            p3.font.name = "Segoe UI"
            p3.font.size = Pt(10.5)
            p3.font.color.rgb = COLOR_TEXT_MUTED

        # -------------------------------------------------------------
        # Slide 9: Diagram 1 - System Use Case Diagram
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="System Use Case & Actor Interaction Model",
            image_filename="use_case_diagram.png",
            category="SECTION 2: UML USE CASE DIAGRAM",
            takeaways=[
                ("Primary Actors", "4 Human actors: Guest Patron, Registered Member, Librarian Staff, and Administrator."),
                ("External Gateways", "Safaricom Daraja (M-Pesa STK Push), OpenAI/Gemini (Catalog RAG), and Open Library (ISBN Import)."),
                ("Member Services", "Physical borrowing, FIFO hold queues, in-app digital reader, M-Pesa pass subscriptions, and AI chatbot."),
                ("Staff Operations", "Circulation desk checkout/return, fine waivers, hold approvals, and catalog inventory CRUD."),
                ("Administrative Core", "User ban management, librarian creation, dynamic membership tiers, and encrypted AI settings.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 10: Technologies Matrix - Core Stack (Table Slide)
        # -------------------------------------------------------------
        self.add_table_slide(
            title="Technology Stack Matrix: Core Architecture",
            category="SECTION 3: TECHNOLOGIES & FRAMEWORKS",
            headers=["Layer", "Technology", "Version", "Key Responsibilities in MaktabaBora"],
            rows=[
                ["Backend Runtime", "PHP", "8.2+", "Type-safe runtime, native vector math, OpenSSL cryptographic functions"],
                ["Backend Framework", "Laravel Framework", "11.x", "Headless RESTful API, Eloquent ORM, IoC Service Container, Gates"],
                ["Dependency Mgr", "Composer", "2.x", "PHP package dependency resolution and class autoloader optimization"],
                ["Relational Database", "PostgreSQL", "16.x", "3NF ACID storage, JSONB indexing, foreign key constraints, cascading"],
                ["Serverless Cloud DB", "Neon Cloud PostgreSQL", "Active", "Cloud DB with pooled endpoints & automatic migration endpoint switching"],
                ["Testing Database", "SQLite (:memory:)", "3.x", "Isolated in-RAM database executing 127 Unit/Feature tests in ~7 seconds"],
                ["Frontend Framework", "React", "18.3.1", "Single-Page Application, custom Context state (AuthContext, CartContext)"],
                ["Build & Dev Server", "Vite", "7.0.7", "Ultra-fast frontend build tooling and Hot Module Replacement (HMR)"]
            ],
            col_widths=[1.8, 2.2, 1.2, 6.533]
        )

        # -------------------------------------------------------------
        # Slide 11: Technologies Matrix - Integrations & QA (Table Slide)
        # -------------------------------------------------------------
        self.add_table_slide(
            title="Technology Stack Matrix: Integrations & Quality",
            category="SECTION 3: TECHNOLOGIES & FRAMEWORKS",
            headers=["Subsystem", "Tool / Service", "Version / Model", "Integration Purpose & Functionality"],
            rows=[
                ["Client Styling", "Tailwind CSS", "v4.0.0", "Utility-first design engine powering responsive layouts & reader themes"],
                ["UI Components", "Lucide React", "0.400.0", "Modern iconography suite across all patron, staff, and admin dashboards"],
                ["UI Motion", "Framer Motion", "12.43.0", "Fluid animations for drawers, reader controls, and checkout countdowns"],
                ["HTTP Client", "Axios", "1.11.0", "Client API library with request interceptors attaching Bearer JWT tokens"],
                ["Payment Gateway", "Safaricom Daraja API", "v2.0", "Customer-to-Business STK Push and real-time transaction polling"],
                ["Catalog AI RAG", "OpenAI / Gemini API", "gpt-4o-mini / flash", "Text embeddings (1536-dim) & grounded conversational AI guidance"],
                ["Metadata Provider", "Open Library API", "REST v1", "1-Click automated ISBN metadata search and cover synchronization"],
                ["Automated Testing", "PHPUnit", "11.x", "Unit test suite (80 tests) and Feature HTTP pipeline suite (47 tests)"]
            ],
            col_widths=[1.8, 2.2, 1.6, 6.133]
        )

        # -------------------------------------------------------------
        # Slide 12: Database Design Principles (Light Theme)
        # -------------------------------------------------------------
        s12 = self.add_blank_slide(is_dark=False)
        self.add_header(s12, "Relational Database Design Principles", category="SECTION 4: DATABASE STRUCTURE", is_dark=False)

        principles = [
            ("Third Normal Form (3NF)", "Zero transitive dependencies. Tables store discrete, non-redundant domain entities with clear single-responsibility boundaries."),
            ("Strict Referential Integrity", "Foreign key constraints with ON DELETE CASCADE for dependent records (e.g. book copies) and ON DELETE RESTRICT for critical financial loans."),
            ("High-Performance Indexing", "Dedicated B-tree indexes on authentication emails, barcodes, ISBNs, loan due dates, and AI usage timestamps for instant querying."),
            ("ACID Financial Compliance", "Multi-table state mutations (e.g. loan issuance, hold fulfillments, pass activations) run inside strict database transactions (DB::beginTransaction)."),
            ("JSONB Semi-Structured Data", "Utilizes native PostgreSQL JSONB columns for digital eBook chapter indexes, reader table-of-contents, and AI grounding metadata."),
            ("Serverless Pooler Awareness", "Backend dynamically switches between PgBouncer pooled connections (app queries) and direct connections (schema migrations).")
        ]

        for i, (p_title, p_desc) in enumerate(principles):
            c_left = 0.8 + (i % 3) * 3.95
            c_top = 1.6 + (i // 3) * 2.75
            card = self.add_card(s12, c_left, c_top, 3.8, 2.5)

            tb_p = s12.shapes.add_textbox(Inches(c_left + 0.25), Inches(c_top + 0.25), Inches(3.3), Inches(2.0))
            tf_p = tb_p.text_frame
            tf_p.word_wrap = True
            tf_p.margin_left = tf_p.margin_top = tf_p.margin_right = tf_p.margin_bottom = 0

            p1 = tf_p.paragraphs[0]
            p1.text = p_title
            p1.font.name = "Segoe UI"
            p1.font.size = Pt(13)
            p1.font.bold = True
            p1.font.color.rgb = COLOR_ACCENT_BLUE
            p1.space_after = Pt(8)

            p2 = tf_p.add_paragraph()
            p2.text = p_desc
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(11)
            p2.font.color.rgb = COLOR_TEXT_PRIMARY

        # -------------------------------------------------------------
        # Slide 13: Diagram 2 - Database Schema ER Diagram
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="Relational Database Schema & ER Diagram",
            image_filename="database_schema_er_diagram.png",
            category="SECTION 4: DATABASE STRUCTURE",
            takeaways=[
                ("Core Entities", "17 Relational tables capturing Identity, Circulation, Digital Media, Subscriptions, AI, and Auditing."),
                ("Identity Model", "Polymorphic extension: users (1) extends to members (patrons) and librarians (staff)."),
                ("Circulation Triad", "books (catalog) contains book_copies (physical barcodes) which are checked out in loans."),
                ("Financial Audits", "Fines linked to loans; Subscriptions linked to tiers; Refund requests audited by librarians."),
                ("AI Memory & Logs", "chat_sessions store chat_messages; ai_usage_logs track rolling 24-hour token consumption.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 14: Data Dictionary - Users & Tiers (Table Slide)
        # -------------------------------------------------------------
        self.add_table_slide(
            title="Data Dictionary: Identity, Patrons & Membership Tiers",
            category="SECTION 4: DATABASE SCHEMA DICTIONARY",
            headers=["Table", "Column", "Data Type", "Constraints", "Description & Operational Rules"],
            rows=[
                ["users", "id", "BIGINT", "PK, Auto-Inc", "Unique global user identity primary key"],
                ["users", "email", "VARCHAR(255)", "UNIQUE, Indexed", "User login address; validated against RFC standards"],
                ["users", "role", "ENUM", "admin, librarian, member", "Primary RBAC gate attribute (Default: 'member')"],
                ["members", "member_number", "VARCHAR(64)", "UNIQUE, Indexed", "Patron system barcode identifier (e.g. MEM-4F8A9B)"],
                ["members", "borrow_limit", "INT", "Default: 3", "Max concurrent active physical book loans allowed"],
                ["members", "is_banned", "BOOLEAN", "Default: false", "Suspension flag; checked by CheckBannedStatus middleware"],
                ["librarians", "employee_number", "VARCHAR(64)", "UNIQUE, Indexed", "Official staff identifier assigned upon creation"],
                ["membership_tiers", "price", "DECIMAL(10,2)", "NOT NULL", "Monthly subscription price in Kenya Shillings (KES)"],
                ["membership_tiers", "digital_discount_pct", "DECIMAL(5,2)", "Default: 0.00", "Percentage discount applied to digital eBook purchases in cart"]
            ],
            col_widths=[1.6, 1.8, 1.6, 2.0, 4.733]
        )

        # -------------------------------------------------------------
        # Slide 15: Data Dictionary - Inventory & Loans (Table Slide)
        # -------------------------------------------------------------
        self.add_table_slide(
            title="Data Dictionary: Inventory, Barcodes & Loans",
            category="SECTION 4: DATABASE SCHEMA DICTIONARY",
            headers=["Table", "Column", "Data Type", "Constraints", "Description & Operational Rules"],
            rows=[
                ["books", "isbn", "VARCHAR(32)", "UNIQUE, Indexed", "10 or 13-digit International Standard Book Number"],
                ["books", "available_copies", "INT", "Default: 1", "Current physical volumes on shelf; decrements on loan"],
                ["books", "rental_price", "DECIMAL(10,2)", "Default: 0.00", "Baseline digital eBook purchase price in KES"],
                ["book_copies", "barcode", "VARCHAR(64)", "UNIQUE, Indexed", "Physical copy barcode scanned at circulation desk"],
                ["book_copies", "status", "ENUM", "avail, borrowed, lost...", "State of physical volume (Default: 'available')"],
                ["loans", "due_date", "TIMESTAMP", "Indexed", "Return deadline (Default: loan_date + 14 days)"],
                ["loans", "status", "ENUM", "active, returned, overdue", "Circulation state; checked by cron/middleware"],
                ["fines", "amount", "DECIMAL(10,2)", "NOT NULL", "Accrued overdue fine (calculated at KES 10.00 / day)"],
                ["reservations", "queue_position", "INT", "NOT NULL", "FIFO waitlist rank for checked-out book (1, 2, 3...)"]
            ],
            col_widths=[1.6, 1.8, 1.6, 2.0, 4.733]
        )

        # -------------------------------------------------------------
        # Slide 16: Data Dictionary - Digital & AI (Table Slide)
        # -------------------------------------------------------------
        self.add_table_slide(
            title="Data Dictionary: Digital Store, AI Memory & Auditing",
            category="SECTION 4: DATABASE SCHEMA DICTIONARY",
            headers=["Table", "Column", "Data Type", "Constraints", "Description & Operational Rules"],
            rows=[
                ["digital_books", "file_format", "ENUM", "'epub', 'pdf'", "Digital eBook MIME format served to in-browser reader"],
                ["digital_books", "chapter_index", "JSONB", "Nullable", "Structured chapter list, anchors, and word counts"],
                ["digital_rentals", "access_type", "ENUM", "'lifetime_purchase'", "Permanent digital license for reading purchased eBook"],
                ["subscriptions", "transaction_reference", "VARCHAR(128)", "Nullable", "Safaricom Daraja M-Pesa receipt verification code"],
                ["chat_sessions", "session_token", "VARCHAR(64)", "UNIQUE, Indexed", "Cryptographic token isolating patron conversation context"],
                ["chat_messages", "grounding_metadata", "JSONB", "Nullable", "Referenced catalog book IDs, authors, and similarity scores"],
                ["ai_usage_logs", "tokens_consumed", "INT", "NOT NULL", "Prompt + completion tokens used for daily quota tracking"],
                ["refund_requests", "status", "ENUM", "pending, approved...", "Formal patron reimbursement workflow audited by staff"]
            ],
            col_widths=[1.6, 1.8, 1.6, 2.0, 4.733]
        )

        # -------------------------------------------------------------
        # Slide 17: Application Directory Overview (Light Theme)
        # -------------------------------------------------------------
        s17 = self.add_blank_slide(is_dark=False)
        self.add_header(s17, "Application Structure & Modular Directory Layout", category="SECTION 5: SYSTEM STRUCTURE", is_dark=False)

        dirs = [
            ("Backend Controllers (`app/Http/Controllers`)", "18 REST controllers managing Auth, Physical Inventory, Loans, Fines, Reservations, Subscriptions, AI Chatbot, and Admin CRUD. Formats input and responses."),
            ("Security Middleware (`app/Http/Middleware`)", "20 Custom gatekeepers enforcing JWT validation, CORS, token-bucket rate limits, borrow quotas, banned member blocks, and daily 20k AI token cost limits."),
            ("Domain Services (`app/Services`)", "19 Decoupled business logic classes implementing explicit interfaces. Covers JWT signing, Daraja STK push, vector cosine scoring, and Open Library imports."),
            ("Frontend Single Page App (`frontend/src`)", "13 Page views (`pages/`), reusable components (`components/`), React Context managers (`AuthContext`, `CartContext`), and Axios API service interceptors.")
        ]

        for i, (d_title, d_desc) in enumerate(dirs):
            c_left = 0.8 + (i % 2) * 5.95
            c_top = 1.6 + (i // 2) * 2.7
            self.add_card(s17, c_left, c_top, 5.75, 2.5, title=d_title)

            tb_d = s17.shapes.add_textbox(Inches(c_left + 0.3), Inches(c_top + 0.65), Inches(5.15), Inches(1.7))
            tf_d = tb_d.text_frame
            tf_d.word_wrap = True
            tf_d.margin_left = tf_d.margin_top = tf_d.margin_right = tf_d.margin_bottom = 0

            p = tf_d.paragraphs[0]
            p.text = d_desc
            p.font.name = "Segoe UI"
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_TEXT_PRIMARY

        # -------------------------------------------------------------
        # Slide 18: Diagram 3 - Eloquent Models Architecture
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="Eloquent Models Architecture & Relationships",
            image_filename="class_diagram_01_models.png",
            category="SECTION 5: CLASS DIAGRAMS",
            takeaways=[
                ("Entity Classes", "14 Eloquent models encapsulating database attributes, mutators, and relationships."),
                ("Identity Hierarchy", "User hasOne Member and Librarian; Member owns Loans, Holds, Passes, and AI Logs."),
                ("Catalog Composition", "Book hasMany BookCopies; Book owns DigitalBook; BookCopy hasMany Loans."),
                ("Financial Links", "Loan hasOne Fine; Subscription belongsTo MembershipTier and hasOne RefundRequest."),
                ("AI Chat Memory", "ChatSession hasMany ChatMessages with grounding metadata and member ownership.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 19: Diagram 4 - REST Controllers Layer
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="REST Controllers Layer Architecture",
            image_filename="class_diagram_02_controllers.png",
            category="SECTION 5: CLASS DIAGRAMS",
            takeaways=[
                ("Subsystem 1: Auth & Patron", "AuthController, LoanController, FineController, ReservationController."),
                ("Subsystem 2: Digital & Media", "DigitalRentalController, SubscriptionController, MembershipTierController."),
                ("Subsystem 3: AI & Discovery", "AiChatbotController, BookRecommendationController, CatalogSearchController."),
                ("Subsystem 4: Staff & Admin", "LibrarianDashboardController, AdminDashboardController, AdminCrudController."),
                ("Architectural Rule", "Controllers are thin HTTP adapters; all domain tasks delegate to injected Services.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 20: Diagram 5 - Domain Services & Contracts
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="Domain Services & Contracts Layer Architecture",
            image_filename="class_diagram_03_services.png",
            category="SECTION 5: CLASS DIAGRAMS",
            takeaways=[
                ("Interface Decoupling", "All services implement strict PHP Contracts in app/Contracts/Services/."),
                ("IoC Container", "Bound as singletons in Service Providers (e.g. AuthServiceProvider)."),
                ("Core Domain Classes", "AuthSessionService, DarajaPaymentService, DigitalRentalService, BookAvailabilityService."),
                ("Algorithms & Machine Learning", "BookEmbeddingService (Cosine vector distance), CatalogRetrievalService (TF-IDF)."),
                ("External Integration", "OpenLibraryService (1-Click ISBN import), CurrencyConverterService (Forex parity).")
            ]
        )

        # -------------------------------------------------------------
        # Slide 21: Diagram 6 - Workflow 1: Authentication & JWT
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="Workflow 1: Authentication & JWT Token Lifecycle",
            image_filename="sequence_01_auth_jwt_lifecycle.png",
            category="SECTION 6: OPERATIONAL WORKFLOWS",
            takeaways=[
                ("Stateless Architecture", "Tokens signed using HMAC-SHA256 with server master key (APP_KEY); zero DB sessions."),
                ("Expiration Policies", "Standard access tokens: 15 minutes; Persistent 'Remember Me' tokens: 30 days."),
                ("Middleware Checkpoint", "JwtTokenValidation inspects Bearer header and sets UserResolver on the request."),
                ("Session Revocation", "Logout explicitly writes token hash to Redis/Cache blacklist with remaining TTL."),
                ("First-Login Rotation", "Librarians are forced to rotate temporary passwords before accessing portal.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 22: Diagram 7 - Workflow 2: Circulation, Returns & Fines
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="Workflow 2: Physical Circulation, Returns & Overdue Fines",
            image_filename="sequence_02_physical_circulation_fines.png",
            category="SECTION 6: OPERATIONAL WORKFLOWS",
            takeaways=[
                ("Pre-Checkout Guards", "ValidateBorrowLimit and CheckFineAmount verify patron eligibility before checkout."),
                ("Atomic Checkout", "DB transaction updates BookCopy status = 'borrowed' and decrements available copies."),
                ("Automated Overdue Calculation", "System computes difference between return date and due date at KES 10.00 / day."),
                ("Fine Settlement Options", "Patron can pay online via Safaricom Daraja M-Pesa STK push."),
                ("Staff Fine Waivers", "Librarians can waive fines with administrative justification and audit logging.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 23: Diagram 8 - Workflow 3: Hold Reservation Queue
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="Workflow 3: Hold Reservation Queue Lifecycle",
            image_filename="sequence_03_hold_reservation_queue.png",
            category="SECTION 6: OPERATIONAL WORKFLOWS",
            takeaways=[
                ("FIFO Queue Allocation", "Patrons join waitlist when available_copies == 0; assigned position count + 1."),
                ("Automated Reordering", "When a patron cancels or copy is fulfilled, all subsequent queue positions decrement."),
                ("Staff Approval Desk", "Librarian approves next in line when copy is returned; allocates 3-day pickup window."),
                ("Printable Receipts", "Patrons receive an interactive, printable digital reservation receipt with barcode."),
                ("Expiration Cleanup", "Holds not picked up within the 3-day window are automatically expired.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 24: Diagram 9 - Workflow 4: Digital Purchase & Reader
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="Workflow 4: Digital Book Purchase & In-App Reader",
            image_filename="sequence_04_digital_purchase_mpesa_reader.png",
            category="SECTION 6: OPERATIONAL WORKFLOWS",
            takeaways=[
                ("Multi-Item Shopping Cart", "Patrons batch-purchase digital eBooks with automatic membership pass discounts."),
                ("Safaricom STK Prompt", "DarajaPaymentService initiates STK push directly to patron's mobile phone."),
                ("Lifetime License", "Confirmed payment mints a DigitalRental record with access_type = 'lifetime_purchase'."),
                ("Protected Reading Stream", "EnsureValidDigitalAccess middleware verifies active purchase before file streaming."),
                ("In-Browser Reader", "Supports EPUB & PDF reading, Light/Dark/Sepia themes, font scaling, and local notes.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 25: Diagram 10 - Workflow 5: Membership Passes & Refunds
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="Workflow 5: Membership Perk Passes & Refund Flow",
            image_filename="sequence_05_membership_perks_refund.png",
            category="SECTION 6: OPERATIONAL WORKFLOWS",
            takeaways=[
                ("Pass Subscription", "Patron subscribes to General, Student, Scholar, or VIP tier via M-Pesa STK push."),
                ("Auto-Upgrade Handling", "Purchasing a new tier automatically retires and cancels any previous active pass."),
                ("Perk Activation", "Instantly updates member borrow limits (up to 10 books) and cart discount rates."),
                ("Dispute Filing", "Patron submits formal refund request with justification via member profile."),
                ("Staff Audit & Resolution", "Librarian reviews claim; approving refund cancels pass and resets borrow quotas.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 26: Diagram 11 - Workflow 6: AI Librarian Assistant (RAG)
        # -------------------------------------------------------------
        self.add_diagram_slide(
            title="Workflow 6: AI Librarian Assistant (Rate-Limited RAG)",
            image_filename="sequence_06_ai_librarian_rag_assistant.png",
            category="SECTION 6: OPERATIONAL WORKFLOWS",
            takeaways=[
                ("Daily Cost Limiter", "ChatbotCostLimiter middleware checks cumulative 24h tokens; halts at 20,000 ceiling."),
                ("Semantic Embeddings", "Query vectorized via text-embedding-3-small and matched using cosine similarity."),
                ("Catalog Context Injection", "System prompt is augmented with real library holdings, authors, and shelf codes."),
                ("Anti-Hallucination Defense", "Explicit system instructions bar LLM from fabricating non-existent books."),
                ("Token Consumption Logging", "Prompt and completion tokens logged to ai_usage_logs on every turn.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 27: Key Code Highlights - JWT & AI Quota (Code Slide)
        # -------------------------------------------------------------
        self.add_code_slide(
            title="Key Code: Stateless JWT & 24h AI Cost Limiter",
            category="SECTION 8: IMPLEMENTATION HIGHLIGHTS",
            code_text="""// 1. Stateless JWT Generation (AuthSessionService.php)
public function generateToken(User $user, bool $remember = false): string {
    $ttl = $remember ? (60 * 60 * 24 * 30) : (60 * 15); // 30d vs 15m
    $header = base64_encode(json_encode(['typ' => 'JWT', 'alg' => 'HS256']));
    $payload = base64_encode(json_encode([
        'iss' => 'MaktabaBora', 'sub' => $user->id,
        'role' => $user->role, 'iat' => time(), 'exp' => time() + $ttl
    ]));
    $sig = hash_hmac('sha256', "{$header}.{$payload}", $this->secret, true);
    return "{$header}.{$payload}." . base64_encode($sig);
}

// 2. Rolling 24h AI Quota Enforcement (ChatbotCostLimiter.php)
public function handle(Request $request, Closure $next): Response {
    $user = $request->user();
    if ($user && ($member = Member::where('user_id', $user->id)->first())) {
        $recentTokens = AiUsageLog::where('member_id', $member->id)
            ->where('created_at', '>=', now()->subDay())->sum('tokens_consumed');
        if ($recentTokens >= 20000) {
            return response()->json(['error' => 'AI Quota Exceeded (20k)'], 429);
        }
    }
    return $next($request);
}""",
            explanation_points=[
                ("Stateless Authentication", "JWTs encode user ID and RBAC role, cryptographically signed with HMAC-SHA256. Avoids costly database lookups on routine requests."),
                ("Cache Blacklisting", "Logout immediately writes the token hash to Redis/Cache with a TTL matching the token's remaining lifetime."),
                ("AI Cost Ceiling", "ChatbotCostLimiter aggregates token consumption over a rolling 24-hour window, halting requests with HTTP 429 before expensive LLM API invocations.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 28: Key Code Highlights - Vector Cosine & TF-IDF (Code Slide)
        # -------------------------------------------------------------
        self.add_code_slide(
            title="Key Code: Vector Cosine Similarity & TF-IDF",
            category="SECTION 8: IMPLEMENTATION HIGHLIGHTS",
            code_text="""// 1. Native High-Dimensional Cosine Similarity (BookEmbeddingService.php)
public static function cosine(array $a, array $b): float {
    $dot = 0.0; $normA = 0.0; $normB = 0.0;
    $len = count($a);
    if ($len === 0 || $len !== count($b)) return 0.0;

    for ($i = 0; $i < $len; $i++) {
        $dot += $a[$i] * $b[$i];
        $normA += $a[$i] * $a[$i];
        $normB += $b[$i] * $b[$i];
    }
    $denom = sqrt($normA) * sqrt($normB);
    return $denom > 0.0 ? ($dot / $denom) : 0.0;
}

// 2. Content-Based TF-IDF Cosine Ranking (CatalogRetrievalService.php)
public function searchCatalog(string $query, ?string $genre = null): Collection {
    $books = Book::query()->when($genre, fn($q) => $q->where('genre', $genre))->get();
    return $books->map(function ($book) use ($query) {
        $book->similarity_score = TextSimilarity::cosineSimilarity(
            $query, "{$book->title} {$book->author} {$book->description}"
        );
        return $book;
    })->filter(fn($b) => $b->similarity_score > 0.15)
      ->sortByDesc('similarity_score')->values();
}""",
            explanation_points=[
                ("Pure PHP Vector Math", "Computes dot product and L2 norms between 1536-dimensional embedding vectors without requiring specialized PostgreSQL vector extensions."),
                ("Hybrid Recommendation", "Combines semantic vector embeddings with term-frequency inverse-document-frequency (TF-IDF) string distance as an offline fallback."),
                ("Grounded Catalog Context", "Matched titles, authors, and shelf codes are extracted and injected into the LLM system prompt to prevent hallucination.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 29: Key Code Highlights - Daraja STK & Forex Caching (Code Slide)
        # -------------------------------------------------------------
        self.add_code_slide(
            title="Key Code: Safaricom Daraja M-Pesa & Forex Caching",
            category="SECTION 8: IMPLEMENTATION HIGHLIGHTS",
            code_text="""// 1. M-Pesa STK Push Prompt Dispatch (DarajaPaymentService.php)
public function initiateStkPush(string $phone, float $amount, string $ref): array {
    $timestamp = date('YmdHis');
    $password = base64_encode($this->shortcode . $this->passkey . $timestamp);

    $res = Http::withToken($this->generateAccessToken())
        ->post('https://sandbox.safaricom.co.ke/mpesa/stkpush/v1/processrequest', [
            'BusinessShortCode' => $this->shortcode,
            'Password'          => $password,
            'Timestamp'         => $timestamp,
            'TransactionType'   => 'CustomerPayBillOnline',
            'Amount'            => (int) ceil($amount),
            'PartyA'            => $this->formatPhoneNumber($phone),
            'PartyB'            => $this->shortcode,
            'PhoneNumber'       => $this->formatPhoneNumber($phone),
            'CallBackURL'       => config('services.daraja.callback_url'),
            'AccountReference'  => substr($ref, 0, 12),
            'TransactionDesc'   => 'MaktabaBora Pay',
        ]);
    return $res->json() ?? [];
}

// 2. 24-Hour Currency Exchange Rate Caching (CurrencyConverterService.php)
public function convert(?float $amount, ?string $from, string $to = 'KES'): ?float {
    if (!$amount || !$from || strtoupper($from) === strtoupper($to)) return $amount;
    $rates = Cache::remember('forex_rates_kes', 86400, fn() => $this->fetchLiveRates());
    if (!isset($rates[$from]) || !isset($rates[$to])) return null;
    return round(($amount / $rates[$from]) * $rates[$to], 2);
}""",
            explanation_points=[
                ("Base64 Security Token", "Generates dynamic Daraja timestamp password and bearer token for authenticated Safaricom gateway communication."),
                ("Simulated Evaluation Flow", "Constructs compliant Daraja 2.0 payload envelopes and mock settlements for safe zero-cost MVP demonstration."),
                ("24-Hour Cache Layer", "Live exchange rates are cached for 86,400 seconds (24h) to avoid third-party API latency and rate-limit penalties.")
            ]
        )

        # -------------------------------------------------------------
        # Slide 30: RESTful API Routing & Gateway (Table Slide)
        # -------------------------------------------------------------
        self.add_table_slide(
            title="REST API Architecture & Public / Member Routes",
            category="SECTION 9: APIS & INTEGRATIONS",
            headers=["Verb", "Endpoint URI", "Controller Action", "Middleware Guards & Authorization"],
            rows=[
                ["POST", "/api/v1/auth/register", "AuthController@register", "Public: Validates unique email, creates member"],
                ["POST", "/api/v1/auth/login", "AuthController@login", "Public: Verifies credentials, mints Bearer JWT"],
                ["GET", "/api/v1/books", "BookInventoryController@index", "Public: Paginated catalog with availability & pricing"],
                ["GET", "/api/v1/books/{book}", "BookInventoryController@show", "Public: Full synopsis, physical shelf code, copies"],
                ["POST", "/api/v1/fines/daraja/callback", "FineController@darajaCallback", "Public: Safaricom M-Pesa webhook callback receiver"],
                ["GET", "/api/v1/auth/me", "AuthController@me", "jwt.validation, ensure.account: Returns user & profile"],
                ["GET", "/api/v1/loans", "LoanController@index", "ensure.member: List member active and historical loans"],
                ["POST", "/api/v1/digital-books/checkout-cart", "DigitalRentalController@checkoutCart", "ensure.member: M-Pesa STK push for eBook cart checkout"],
                ["GET", "/api/v1/digital-books/{id}/read", "DigitalRentalController@read", "ensure.digital_access: Streams chapter index & reader text"]
            ],
            col_widths=[1.4, 3.8, 2.8, 3.733]
        )

        # -------------------------------------------------------------
        # Slide 31: Staff & Admin Endpoints (Table Slide)
        # -------------------------------------------------------------
        self.add_table_slide(
            title="REST API: Circulation Desk & Administration Routes",
            category="SECTION 9: APIS & INTEGRATIONS",
            headers=["Verb", "Endpoint URI", "Controller Action", "Middleware Guards & Authorization"],
            rows=[
                ["POST", "/api/v1/ai/chat", "AiChatbotController@chat", "chatbot.cost_limiter: Grounded RAG conversational AI"],
                ["GET", "/api/v1/librarian/metrics", "LibrarianDashboardController@metrics", "ensure.librarian: Circulation statistics, active loans, fines"],
                ["POST", "/api/v1/librarian/loans/checkout", "LoanController@checkout", "validate.borrow_limit, check.fine: Barcode scan checkout"],
                ["POST", "/api/v1/librarian/loans/{id}/return", "LoanController@returnBook", "ensure.librarian: Check-in volume & auto-accrue fines"],
                ["POST", "/api/v1/librarian/fines/{id}/waive", "FineController@waive", "ensure.librarian: Waive overdue fine with audit trail"],
                ["POST", "/api/v1/librarian/openlibrary/import", "LibrarianDashboardController@import", "ensure.librarian: 1-Click Open Library ISBN catalog import"],
                ["GET", "/api/v1/admin/users", "AdminDashboardController@users", "ensure.admin: List system users, roles, and ban statuses"],
                ["POST", "/api/v1/admin/users/{user}/toggle-ban", "AdminDashboardController@toggleBan", "ensure.admin: Suspend user account with mandatory reason"],
                ["GET", "/api/v1/admin/crud/{table}", "AdminCrudController@getTableData", "ensure.admin: Direct administrative table maintenance"]
            ],
            col_widths=[1.4, 3.8, 3.0, 3.533]
        )

        # -------------------------------------------------------------
        # Slide 32: Security & Defense-in-Depth (Light Theme)
        # -------------------------------------------------------------
        s32 = self.add_blank_slide(is_dark=False)
        self.add_header(s32, "Security Architecture & Defense-in-Depth", category="SECTION 10: SECURITY CONSIDERATIONS", is_dark=False)

        sec_pillars = [
            ("Stateless Cryptographic Tokens", "HMAC-SHA256 JWT tokens signed with master secret. Payload tampering invalidates signatures instantly. Logout adds token hashes to Redis cache blacklist with automatic TTL eviction."),
            ("Sequential Middleware Gatekeepers", "Requests traverse an ordered inspection pipeline: CORS -> Rate Limiting -> JWT Verification -> Account Active Check -> Non-Banned Check -> Borrow Limit -> Fine Check."),
            ("Financial Transaction Integrity", "M-Pesa STK callbacks verify merchant shortcodes and transaction IDs. Multi-table asset grants and fine clearances execute inside ACID database transactions to prevent race conditions."),
            ("AI Anti-Hallucination & PII Safety", "RAG prompts strictly constrain LLM context to real database catalog holdings. Explicit system prompts bar the assistant from discussing user accounts or disclosing patron borrowing history."),
            ("SQL Injection Defense", "All database read and write operations are routed through Eloquent ORM or parameterized PDO prepared statements, fully immunizing the platform from SQL injection attacks."),
            ("Credential & API Key Encryption", "External third-party API keys (OpenAI, Gemini, Daraja) are stored encrypted using OpenSSL AES-256-CBC and masked when displayed in administrative interfaces.")
        ]

        for i, (s_title, s_desc) in enumerate(sec_pillars):
            c_left = 0.8 + (i % 3) * 3.95
            c_top = 1.6 + (i // 3) * 2.75
            card = self.add_card(s32, c_left, c_top, 3.8, 2.5)

            tb_s = s32.shapes.add_textbox(Inches(c_left + 0.25), Inches(c_top + 0.25), Inches(3.3), Inches(2.0))
            tf_s = tb_s.text_frame
            tf_s.word_wrap = True
            tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0

            p1 = tf_s.paragraphs[0]
            p1.text = s_title
            p1.font.name = "Segoe UI"
            p1.font.size = Pt(13)
            p1.font.bold = True
            p1.font.color.rgb = COLOR_ACCENT_BLUE
            p1.space_after = Pt(8)

            p2 = tf_s.add_paragraph()
            p2.text = s_desc
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(11)
            p2.font.color.rgb = COLOR_TEXT_PRIMARY

        # -------------------------------------------------------------
        # Slide 33: Challenges & Engineering Solutions (Light Theme)
        # -------------------------------------------------------------
        s33 = self.add_blank_slide(is_dark=False)
        self.add_header(s33, "Engineering Challenges & Implemented Solutions", category="SECTION 11: CHALLENGES & SOLUTIONS", is_dark=False)

        challenges = [
            ("Challenge 1: Neon Connection Pooling vs DDL", "Neon's PgBouncer transaction pooling aborts DDL schema migrations.", "Solution: Implemented an automatic detector in config/database.php that strips the -pooler. suffix during artisan migrate commands to connect directly."),
            ("Challenge 2: PostgreSQL SSL Mode Mismatch", "Neon requires sslmode=require, while local Windows Postgres throws SQLSTATE[08006].", "Solution: Configured sslmode fallback to 'prefer', allowing local connections without SSL while negotiating SSL with Neon Cloud seamlessly."),
            ("Challenge 3: Unbounded AI API Cost Exposure", "Uncontrolled patron AI queries risk catastrophic monthly API billing spikes.", "Solution: Authored ChatbotCostLimiter middleware enforcing a hard 20,000 token limit per member across rolling 24-hour periods."),
            ("Challenge 4: Localhost M-Pesa Webhooks", "Safaricom cannot dispatch callbacks to local private IP development environments.", "Solution: Built a hybrid polling mechanism in DarajaPaymentService to query transaction status continuously while the modal is displayed."),
            ("Challenge 5: Vector Search without Vector DB", "Hosting dedicated vector database extensions increases infrastructure overhead.", "Solution: Implemented pure PHP high-dimensional cosine similarity ranking, backed by TF-IDF keyword cosine similarity as an offline fallback.")
        ]

        for i, (c_head, c_prob, c_sol) in enumerate(challenges):
            c_top = 1.55 + i * 1.15
            self.add_card(s33, 0.8, c_top, 11.733, 1.05)

            tb_c = s33.shapes.add_textbox(Inches(1.0), Inches(c_top + 0.12), Inches(11.333), Inches(0.8))
            tf_c = tb_c.text_frame
            tf_c.word_wrap = True
            tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0

            p_ch = tf_c.paragraphs[0]
            p_ch.text = c_head
            p_ch.font.name = "Segoe UI"
            p_ch.font.size = Pt(12)
            p_ch.font.bold = True
            p_ch.font.color.rgb = COLOR_ACCENT_BLUE

            p_cp = tf_c.add_paragraph()
            p_cp.text = f"{c_prob} -> {c_sol}"
            p_cp.font.name = "Segoe UI"
            p_cp.font.size = Pt(10.5)
            p_cp.font.color.rgb = COLOR_TEXT_PRIMARY

        # -------------------------------------------------------------
        # Slide 34: Testing & Quality Assurance (Light Theme)
        # -------------------------------------------------------------
        s34 = self.add_blank_slide(is_dark=False)
        self.add_header(s34, "Testing Suite & Quality Assurance Metrics", category="SECTION 12: TESTING & QA", is_dark=False)

        # 3 Top Metric Cards
        metrics = [
            ("80 / 80 Passed", "Unit Tests (185 Assertions)", COLOR_ACCENT_EMERALD),
            ("47 / 47 Passed", "Feature Tests (212 Assertions)", COLOR_ACCENT_BLUE),
            ("127 Total Tests", "397 Assertions in ~7.5 Seconds", COLOR_ACCENT_CYAN)
        ]

        for i, (m_val, m_lbl, m_col) in enumerate(metrics):
            c_left = 0.8 + i * 4.0
            card = s34.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(1.6), Inches(3.7), Inches(1.4))
            card.fill.solid()
            card.fill.fore_color.rgb = COLOR_BG_DARK
            card.line.fill.background()

            tb_m = s34.shapes.add_textbox(Inches(c_left + 0.2), Inches(1.8), Inches(3.3), Inches(1.0))
            tf_m = tb_m.text_frame
            tf_m.word_wrap = True
            p_v = tf_m.paragraphs[0]
            p_v.text = m_val
            p_v.font.name = "Segoe UI"
            p_v.font.size = Pt(22)
            p_v.font.bold = True
            p_v.font.color.rgb = m_col

            p_l = tf_m.add_paragraph()
            p_l.text = m_lbl
            p_l.font.name = "Segoe UI"
            p_l.font.size = Pt(11)
            p_l.font.color.rgb = COLOR_TEXT_LIGHT

        # Bottom Breakdown Cards
        test_scopes = [
            ("Isolated In-Memory Test DB", [
                "Configured in phpunit.xml via SQLite :memory:",
                "Tests execute in RAM in ~7.5 seconds",
                "Uses RefreshDatabase trait for clean state isolation",
                "Zero risk of mutating local or Neon PostgreSQL data"
            ]),
            ("Unit Test Coverage (`tests/Unit`)", [
                "15 Services tests (JWT, Cosine math, M-Pesa, RAG guardrails)",
                "17 Middleware tests (Rate limiters, borrow caps, token guards)",
                "6 Provider tests (IoC container resolution & singletons)",
                "ContractsIntegrityTest asserting strict interface compliance"
            ]),
            ("Feature Test Coverage (`tests/Feature`)", [
                "Full HTTP pipeline tests (POST /api/v1/auth/login, loans)",
                "Digital eBook purchasing and streaming permission checks",
                "Safaricom M-Pesa webhook callback simulation",
                "Admin user ban toggles and dynamic CRUD operations"
            ])
        ]

        for i, (t_title, t_items) in enumerate(test_scopes):
            c_left = 0.8 + i * 4.0
            self.add_card(s34, c_left, 3.2, 3.7, 3.65, title=t_title)

            tb_ts = s34.shapes.add_textbox(Inches(c_left + 0.25), Inches(3.8), Inches(3.2), Inches(2.8))
            tf_ts = tb_ts.text_frame
            tf_ts.word_wrap = True
            tf_ts.margin_left = tf_ts.margin_top = tf_ts.margin_right = tf_ts.margin_bottom = 0

            for j, item in enumerate(t_items):
                p = tf_ts.add_paragraph() if j > 0 else tf_ts.paragraphs[0]
                p.text = f"•  {item}"
                p.font.name = "Segoe UI"
                p.font.size = Pt(10.5)
                p.font.color.rgb = COLOR_TEXT_PRIMARY
                p.space_after = Pt(4)

        # -------------------------------------------------------------
        # -------------------------------------------------------------
        # Slide 35: Unified Fullstack Container Architecture (Light Theme)
        # -------------------------------------------------------------
        s35 = self.add_blank_slide(is_dark=False)
        self.add_header(s35, "Unified Fullstack Container Architecture", category="SECTION 13: DEPLOYMENT & OPERATIONS", is_dark=False)

        cloud_cards = [
            ("1. Multi-Stage Docker Build", "node:20-alpine + php:8.2-fpm-nginx", "Stage 1 compiles React 18 SPA with Vite into /dist. Stage 2 serves Nginx + PHP-FPM 8.2 and copies static SPA into Laravel public/ folder."),
            ("2. Single Turnkey Service", "Render Blueprint (render.yaml)", "Deploys both frontend UI and backend API as a single web service, eliminating CORS complexity and hardcoded local ports."),
            ("3. Serverless Cloud Database", "Neon PostgreSQL (AWS Ohio)", "High-availability serverless PostgreSQL 16 with PgBouncer connection pooling and dynamic direct-migration fallback detection."),
            ("4. SPA Catch-All Route", "Laravel routes/web.php", "Nginx serves static assets directly; web routes (/, /catalog, /login) are routed to index.html for client-side React Router navigation."),
            ("5. Production Optimizations", "OPcache & Bytecode Caching", "Bytecode caching enabled (PHP_OPCACHE_ENABLE=1) and Composer dependencies bundled with optimized autoloader (--no-dev --prefer-dist)."),
            ("6. Simulated M-Pesa Evaluation", "Compliant Sandbox Payment Flow", "Constructs standard Daraja 2.0 payload envelopes and polling loops in simulation mode for zero-cost, safe evaluation.")
        ]

        for i, (title_p, ver_p, desc_p) in enumerate(cloud_cards):
            c_left = 0.8 + (i % 3) * 3.95
            c_top = 1.6 + (i // 3) * 2.75
            card = self.add_card(s35, c_left, c_top, 3.8, 2.5)

            tb_pr = s35.shapes.add_textbox(Inches(c_left + 0.25), Inches(c_top + 0.25), Inches(3.3), Inches(2.0))
            tf_pr = tb_pr.text_frame
            tf_pr.word_wrap = True
            tf_pr.margin_left = tf_pr.margin_top = tf_pr.margin_right = tf_pr.margin_bottom = 0

            p1 = tf_pr.paragraphs[0]
            p1.text = title_p
            p1.font.name = "Segoe UI"
            p1.font.size = Pt(13)
            p1.font.bold = True
            p1.font.color.rgb = COLOR_ACCENT_BLUE

            p2 = tf_pr.add_paragraph()
            p2.text = ver_p
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(10)
            p2.font.bold = True
            p2.font.color.rgb = COLOR_TEXT_MUTED
            p2.space_after = Pt(6)

            p3 = tf_pr.add_paragraph()
            p3.text = desc_p
            p3.font.name = "Segoe UI"
            p3.font.size = Pt(10.5)
            p3.font.color.rgb = COLOR_TEXT_PRIMARY

        # -------------------------------------------------------------
        # Slide 36: Continuous Integration & Deployment (GitHub Actions)
        # -------------------------------------------------------------
        s36 = self.add_blank_slide(is_dark=False)
        self.add_header(s36, "Continuous Integration & Deployment (CI/CD)", category="SECTION 13: DEPLOYMENT & OPERATIONS", is_dark=False)

        # Left Card: Pipeline Architecture
        self.add_card(s36, 0.8, 1.6, 5.75, 5.3, title="GitHub Actions Automated Pipeline")
        tb_ci_l = s36.shapes.add_textbox(Inches(1.05), Inches(2.25), Inches(5.25), Inches(4.5))
        tf_ci_l = tb_ci_l.text_frame
        tf_ci_l.word_wrap = True
        tf_ci_l.margin_left = tf_ci_l.margin_top = tf_ci_l.margin_right = tf_ci_l.margin_bottom = 0

        ci_steps = [
            ("Stage 1: Multi-Branch Trigger", "Workflow activates on push and pull-request events targeting 'main' and 'kimura' development branches."),
            ("Stage 2: Container Environment Setup", "Spins up ubuntu-latest virtual runner, installs PHP 8.2 and Node 20 runtimes, and restores caches."),
            ("Stage 3: Automated Test & Build Gates", "Executes complete PHPUnit suite (127 tests, 397 assertions) and builds the React SPA bundle with Vite."),
            ("Stage 4: Zero-Downtime Deploy Trigger", "On 100% test passage, curls Render Deploy Hook URL with commit metadata. Broken builds are automatically blocked.")
        ]
        for i, (h, d) in enumerate(ci_steps):
            p = tf_ci_l.add_paragraph() if i > 0 else tf_ci_l.paragraphs[0]
            p.space_after = Pt(10)
            r1 = p.add_run()
            r1.text = f"{h}\n"
            r1.font.name = "Segoe UI"
            r1.font.bold = True
            r1.font.size = Pt(11)
            r1.font.color.rgb = COLOR_ACCENT_BLUE
            r2 = p.add_run()
            r2.text = d
            r2.font.name = "Segoe UI"
            r2.font.size = Pt(10)
            r2.font.color.rgb = COLOR_TEXT_PRIMARY

        # Right Card: Quality Controls
        self.add_card(s36, 6.75, 1.6, 5.75, 5.3, title="Quality Assurance & Security Controls")
        tb_ci_r = s36.shapes.add_textbox(Inches(7.0), Inches(2.25), Inches(5.25), Inches(4.5))
        tf_ci_r = tb_ci_r.text_frame
        tf_ci_r.word_wrap = True
        tf_ci_r.margin_left = tf_ci_r.margin_top = tf_ci_r.margin_right = tf_ci_r.margin_bottom = 0

        ci_controls = [
            ("Strict Quality Gatekeeping", "No commit reaches production without passing every assertion. Prevents regressions in loan quotas, fine calculations, and JWT validation."),
            ("Encrypted Webhook Secrets", "The Render Deploy Hook URL is stored strictly in GitHub Encrypted Repository Secrets (RENDER_DEPLOY_HOOK_URL), preventing credential leaks."),
            ("High-Velocity Execution", "Complete test suite runs in under 15 seconds, providing immediate feedback to developers on pull requests."),
            ("Audit Trail & Traceability", "Every deployment links commit SHA, test execution results, and container build logs directly in GitHub and Render consoles.")
        ]
        for i, (h, d) in enumerate(ci_controls):
            p = tf_ci_r.add_paragraph() if i > 0 else tf_ci_r.paragraphs[0]
            p.space_after = Pt(10)
            r1 = p.add_run()
            r1.text = f"{h}\n"
            r1.font.name = "Segoe UI"
            r1.font.bold = True
            r1.font.size = Pt(11)
            r1.font.color.rgb = COLOR_ACCENT_EMERALD
            r2 = p.add_run()
            r2.text = d
            r2.font.name = "Segoe UI"
            r2.font.size = Pt(10)
            r2.font.color.rgb = COLOR_TEXT_PRIMARY

        # -------------------------------------------------------------
        # Slide 37: MVP Deployment Analysis: Advantages vs. Disadvantages
        # -------------------------------------------------------------
        s37 = self.add_blank_slide(is_dark=False)
        self.add_header(s37, "MVP Deployment: Advantages vs. Disadvantages", category="SECTION 13: DEPLOYMENT & OPERATIONS", is_dark=False)

        # Left Column: Advantages
        adv_title_card = s37.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.75), Inches(0.45))
        adv_title_card.fill.solid()
        adv_title_card.fill.fore_color.rgb = COLOR_ACCENT_EMERALD
        adv_title_card.line.fill.background()
        p_at = adv_title_card.text_frame.paragraphs[0]
        p_at.text = "PROS: ARCHITECTURAL ADVANTAGES"
        p_at.font.name = "Segoe UI"
        p_at.font.size = Pt(11)
        p_at.font.bold = True
        p_at.font.color.rgb = COLOR_CARD_BG

        advantages = [
            ("1. Zero Infrastructure Expenditure ($0/mo)", "Render Free Tier + Neon DB", "Fully functional containerized web service and scalable cloud PostgreSQL database with zero financial hosting costs for MVP evaluation."),
            ("2. Fullstack Single-Origin Simplicity", "Unified Container + Auto SSL", "UI and API share the exact same origin, eliminating CORS issues and hardcoded ports with automated Let's Encrypt SSL."),
            ("3. Automated CI/CD Test Gate", "GitHub Actions Deployment Hook", "127 automated tests must pass before the Render deploy hook is triggered, guaranteeing broken code never reaches production.")
        ]
        for i, (t, sub, desc) in enumerate(advantages):
            top_i = 2.05 + i * 1.7
            self.add_card(s37, 0.8, top_i, 5.75, 1.55)
            tb = s37.shapes.add_textbox(Inches(1.05), Inches(top_i + 0.15), Inches(5.25), Inches(1.25))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p1 = tf.paragraphs[0]
            p1.text = t
            p1.font.name = "Segoe UI"
            p1.font.size = Pt(12)
            p1.font.bold = True
            p1.font.color.rgb = COLOR_ACCENT_EMERALD
            p2 = tf.add_paragraph()
            p2.text = sub
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(9.5)
            p2.font.bold = True
            p2.font.color.rgb = COLOR_TEXT_MUTED
            p2.space_after = Pt(4)
            p3 = tf.add_paragraph()
            p3.text = desc
            p3.font.name = "Segoe UI"
            p3.font.size = Pt(10)
            p3.font.color.rgb = COLOR_TEXT_PRIMARY

        # Right Column: Disadvantages / Trade-offs
        dis_title_card = s37.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(1.5), Inches(5.75), Inches(0.45))
        dis_title_card.fill.solid()
        dis_title_card.fill.fore_color.rgb = COLOR_ACCENT_AMBER
        dis_title_card.line.fill.background()
        p_dt = dis_title_card.text_frame.paragraphs[0]
        p_dt.text = "CONS: MVP TRADE-OFFS & LIMITATIONS"
        p_dt.font.name = "Segoe UI"
        p_dt.font.size = Pt(11)
        p_dt.font.bold = True
        p_dt.font.color.rgb = COLOR_CARD_BG

        disadvantages = [
            ("1. Free-Tier Inactivity Sleep", "15-Min Inactivity Timeout", "Render puts free containers to sleep after 15 minutes of inactivity; the initial wake-up request incurs a 30 to 50-second cold start latency delay."),
            ("2. Compute & Memory Ceiling", "0.1 CPU & 512 MB RAM", "Free tier bounds high concurrency and memory-intensive batch operations (easily upgraded to Starter $7/mo for dedicated CPU and no sleep)."),
            ("3. Ephemeral Disk & Simulated Flow", "Stateless Container Architecture", "Container filesystem resets on redeploy (all state safely stored in Neon DB); Daraja M-Pesa operates in simulated flow for evaluation safety.")
        ]
        for i, (t, sub, desc) in enumerate(disadvantages):
            top_i = 2.05 + i * 1.7
            self.add_card(s37, 6.75, top_i, 5.75, 1.55)
            tb = s37.shapes.add_textbox(Inches(7.0), Inches(top_i + 0.15), Inches(5.25), Inches(1.25))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p1 = tf.paragraphs[0]
            p1.text = t
            p1.font.name = "Segoe UI"
            p1.font.size = Pt(12)
            p1.font.bold = True
            p1.font.color.rgb = COLOR_ACCENT_AMBER
            p2 = tf.add_paragraph()
            p2.text = sub
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(9.5)
            p2.font.bold = True
            p2.font.color.rgb = COLOR_TEXT_MUTED
            p2.space_after = Pt(4)
            p3 = tf.add_paragraph()
            p3.text = desc
            p3.font.name = "Segoe UI"
            p3.font.size = Pt(10)
            p3.font.color.rgb = COLOR_TEXT_PRIMARY

        # -------------------------------------------------------------
        # Slide 38: Alternative Bare-Metal Setup & CLI Pipeline (Light Theme)
        # -------------------------------------------------------------
        s38 = self.add_blank_slide(is_dark=False)
        self.add_header(s38, "Alternative Bare-Metal / Local CLI Setup", category="SECTION 13: DEPLOYMENT & OPERATIONS", is_dark=False)

        steps = [
            ("Step 1: Clone Repository & PHP Dependencies", "git clone https://github.com/lornaarwa/Smart-library-management-system.git\ncd backend && composer install --no-dev --optimize-autoloader"),
            ("Step 2: Environment Configuration & Key Generation", "cp .env.example .env\nphp artisan key:generate  # Generates 32-char master secret for JWT & encryption"),
            ("Step 3: Database Schema Migration & Seeding", "php artisan migrate --force --seed  # Populates baseline catalog, tiers & test users"),
            ("Step 4: Cache Optimization & Configuration Locking", "php artisan config:cache && php artisan route:cache && php artisan view:cache"),
            ("Step 5: Frontend Single Page App Build", "cd ../frontend && npm install && npm run build  # Generates static assets in frontend/dist")
        ]

        for i, (s_title, s_cmd) in enumerate(steps):
            s_top = 1.55 + i * 1.15
            self.add_card(s38, 0.8, s_top, 11.733, 1.05)

            tb_s = s38.shapes.add_textbox(Inches(1.0), Inches(s_top + 0.12), Inches(11.333), Inches(0.8))
            tf_s = tb_s.text_frame
            tf_s.word_wrap = True
            tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0

            p_st = tf_s.paragraphs[0]
            p_st.text = s_title
            p_st.font.name = "Segoe UI"
            p_st.font.size = Pt(12)
            p_st.font.bold = True
            p_st.font.color.rgb = COLOR_ACCENT_BLUE

            p_sc = tf_s.add_paragraph()
            p_sc.text = s_cmd
            p_sc.font.name = "Consolas"
            p_sc.font.size = Pt(9.5)
            p_sc.font.color.rgb = COLOR_TEXT_PRIMARY

        # -------------------------------------------------------------
        # Slide 39: Future Improvements & Strategic Roadmap (Light Theme)
        # -------------------------------------------------------------
        s39 = self.add_blank_slide(is_dark=False)
        self.add_header(s39, "Strategic Roadmap & Future Enhancements", category="SECTION 14: FUTURE ROADMAP", is_dark=False)

        roadmap = [
            ("1. Hardware RFID & Barcode Scanners", "WebUSB & HID Scanner Listeners", "Enable hands-free, high-throughput book checkout and inventory auditing directly from browser workstations without auxiliary client drivers."),
            ("2. Encrypted Offline-First PWA Reader", "Progressive Web App & Web Cryptography", "Cache purchased digital eBooks securely via Service Workers with client-side AES-GCM decryption, enabling reading during internet outages."),
            ("3. Automated SMS via Africa's Talking", "Telephony Gateway Integration", "Dispatch automated SMS notifications to patrons 48 hours prior to loan due dates and immediate alerts when waitlisted hold copies are available."),
            ("4. Multi-Branch Inter-Library Sync", "Multi-Tenant Enterprise Scaling", "Synchronize catalogs and manage inter-library loans across multiple university or county library campuses with central governance.")
        ]

        for i, (r_title, r_sub, r_desc) in enumerate(roadmap):
            c_left = 0.8 + (i % 2) * 5.95
            c_top = 1.6 + (i // 2) * 2.7
            self.add_card(s39, c_left, c_top, 5.75, 2.5)

            tb_r = s39.shapes.add_textbox(Inches(c_left + 0.3), Inches(c_top + 0.25), Inches(5.15), Inches(2.0))
            tf_r = tb_r.text_frame
            tf_r.word_wrap = True
            tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0

            p1 = tf_r.paragraphs[0]
            p1.text = r_title
            p1.font.name = "Segoe UI"
            p1.font.size = Pt(13)
            p1.font.bold = True
            p1.font.color.rgb = COLOR_ACCENT_BLUE

            p2 = tf_r.add_paragraph()
            p2.text = r_sub
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(10.5)
            p2.font.bold = True
            p2.font.color.rgb = COLOR_TEXT_MUTED
            p2.space_after = Pt(6)

            p3 = tf_r.add_paragraph()
            p3.text = r_desc
            p3.font.name = "Segoe UI"
            p3.font.size = Pt(11)
            p3.font.color.rgb = COLOR_TEXT_PRIMARY

        # -------------------------------------------------------------
        # Slide 40: Conclusion & Final Evaluation (Dark Theme)
        # -------------------------------------------------------------
        s40 = self.add_blank_slide(is_dark=True)
        self.add_header(s40, "Conclusion & Architectural Summary", category="SECTION 15: CONCLUSION", is_dark=True)

        c_box = s40.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
        tf_c = c_box.text_frame
        tf_c.word_wrap = True

        p_c1 = tf_c.paragraphs[0]
        p_c1.text = "A Superior, Production-Ready Library Management Ecosystem"
        p_c1.font.name = "Segoe UI"
        p_c1.font.size = Pt(26)
        p_c1.font.bold = True
        p_c1.font.color.rgb = COLOR_ACCENT_CYAN
        p_c1.space_after = Pt(16)

        p_c2 = tf_c.add_paragraph()
        p_c2.text = (
            "MaktabaBora redefines modern library management by uniting physical inventory administration with an advanced "
            "digital ecosystem. By combining a headless, interface-driven Laravel 11 REST API, an intuitive React 18 frontend, "
            "robust PostgreSQL 3NF data modeling, cashless Safaricom Daraja M-Pesa payments (simulated for MVP), and grounded Retrieval-Augmented "
            "Generation (RAG) AI, the platform delivers an enterprise-grade, secure, and production-ready solution."
        )
        p_c2.font.name = "Segoe UI"
        p_c2.font.size = Pt(14)
        p_c2.font.color.rgb = COLOR_TEXT_LIGHT
        p_c2.space_after = Pt(24)

        # 4 Badge Pillars on Conclusion
        c_pillars = [
            "✔ 100% Passing Automated Tests (127/127)",
            "✔ 11 High-Resolution Horizontal A4 UML Diagrams",
            "✔ Zero Database Session Bloat (Stateless HMAC-SHA256 JWT)",
            "✔ Containerized Render & Neon DB Deployment with GitHub Actions CI/CD"
        ]
        for pill_t in c_pillars:
            p_pil = tf_c.add_paragraph()
            p_pil.text = pill_t
            p_pil.font.name = "Segoe UI"
            p_pil.font.size = Pt(13)
            p_pil.font.bold = True
            p_pil.font.color.rgb = COLOR_ACCENT_EMERALD
            p_pil.space_after = Pt(8)

        # Save presentation
        print(f"Saving {self.slide_count} slides to {OUTPUT_PPTX}...")
        self.prs.save(OUTPUT_PPTX)
        file_size_mb = os.path.getsize(OUTPUT_PPTX) / (1024 * 1024)
        print(f"Successfully generated MaktabaBora presentation! ({file_size_mb:.2f} MB, {self.slide_count} slides)")


if __name__ == "__main__":
    builder = MaktabaBoraDeckBuilder()
    builder.build_deck()

